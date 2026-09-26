# -*- coding: utf-8 -*-
"""
Смоук титульного листа «О отчёте» (io_module/report_cover.py).

Проверяет:
  1. Сбор строк для рублёвой компании: заказ, статусы, режимы, допуски, легенда.
  2. Сбор строк для валютной компании с расхождением ЧП = НРП и результатом
     сверки 1c: статус НЕ СВЕДЕНО должен быть явно помечен в файле.
  3. Единый формат допусков: доли -> проценты, суммы -> «тыс.ед.» с разрядами.
  4. Запись в реальный xlsx: «О отчёте» первым листом, цвета ярлычков,
     ширины колонок, перенос текста, выделение заголовков разделов.
  5. Служебные файлы ~$ не попадают в список диагностики.
  6. Сбой при сборе строк не превращается в падение сохранения отчёта.

Запуск: python _smoke_report_cover.py
"""
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).parent))

from config.defaults import format_tolerance_value
from io_module import report_cover as rc

PASSED = 0


def check(condition: bool, message: str) -> None:
    global PASSED
    if condition:
        PASSED += 1
        print(f"  [OK] {message}")
    else:
        print(f"  [FAIL] {message}")


def make_context(**overrides) -> SimpleNamespace:
    """Контекст с минимальным набором атрибутов, которые читает report_cover."""
    base = dict(
        company='РЗК',
        period='8мес2026',
        type_period='8 месяцев',
        segment='Основной',
        currency='RUB',
        balance_date=None,
        name_file_general_osv='РЗК_общаяосв_нд_2025_.xlsx',
        run_id='20260926_171022',
        balance_df=pd.DataFrame({'сальдо, тыс.ед.': [1.0, 2.0]}),
        pnl_df=pd.DataFrame({'Значение': [1.0]}),
        data={},
        tolerance_params={
            'tolerance_balance': 3000.0,
            'tolerance_pnl_balance': 1050.0,
            'tolerance_rate_deviation': 0.3,
            'nds_missing_values': 0.22,
            'tolerance_orphan_distribution': 0.01,
        },
    )
    base.update(overrides)
    return SimpleNamespace(**base)


def find_row(rows: list, needle: str) -> str:
    """Значение строки, в первой колонке которой встречается needle."""
    for label, value in rows:
        if needle in (label or ''):
            return value
    return ''


def find_label(rows: list, needle: str) -> str:
    """Показатель строки, содержащей needle."""
    for label, _ in rows:
        if needle in (label or ''):
            return label
    return ''


def main() -> None:
    tmp_dir = Path(tempfile.mkdtemp(prefix='smoke_cover_'))
    (tmp_dir / 'mismatches').mkdir()
    (tmp_dir / 'warnings').mkdir()
    (tmp_dir / 'mismatches' / 'mismatch_step_19.xlsx').write_bytes(b'x')
    (tmp_dir / 'mismatches' / '~$mismatch_step_19.xlsx').write_bytes(b'x')
    (tmp_dir / 'warnings' / 'nds_missing_rows_РЗК.xlsx').write_bytes(b'x')
    (tmp_dir / 'sort_report.xlsx').write_bytes(b'x')

    # Папка запуска не нужна — подменяем на временную
    rc.get_run_dir = lambda: tmp_dir

    print('\n1. Рублёвая компания, всё сошлось')
    rows = rc.build_cover_rows(make_context(), warnings=[])
    check(find_row(rows, 'Компания') == 'РЗК', 'компания')
    check('8мес2026' in find_row(rows, 'Период отчётности'), 'период с типом периода')
    check('не требуется' in find_row(rows, 'Валюта остатков'), 'RUB без перевода')
    check('СВЕДЕНО' in find_row(rows, 'Увязка ОПУ и баланса'), 'статус увязки ЧП = НРП')
    check(find_row(rows, 'Предупреждений') == 'нет', 'предупреждений нет')
    check(find_row(rows, 'Строк в расшифровке баланса') == '2', 'счётчик строк баланса')
    check('мягкий режим' in find_row(rows, 'Неизвестные контрагенты'), 'режим контрагентов из config')
    check(find_row(rows, 'Сходимость баланса') == '3 000 тыс.ед.', 'допуск баланса с единицей')
    check(find_row(rows, 'Ставка НДС') == '22%', 'ставка НДС в процентах')
    check(find_row(rows, 'Допуск на потерю суммы') == '1%', 'допуск распределения в процентах')
    check(find_row(rows, 'О отчёте').startswith('этот лист'), 'легенда описывает титульный лист')

    print('\n2. Валютная компания, расхождение ЧП = НРП, есть сверка 1c')
    ctx = make_context(
        currency='cny',
        balance_date='31.12.2025',
        data={
            'pnl_balance_mismatch': {'diff': 4321.4, 'tolerance': 1050.0},
            'reconciliation_summary': {'tolerance': 200.0},
        },
    )
    rows = rc.build_cover_rows(ctx, warnings=['a', 'b', 'c'])
    currency_text = find_row(rows, 'Валюта остатков')
    check('CNY' in currency_text and '31.12.2025' in currency_text, f'валюта и дата перевода: {currency_text}')
    mismatch_text = find_row(rows, 'Увязка ОПУ и баланса')
    check('НЕ СВЕДЕНО' in mismatch_text, 'расхождение помечено явно')
    check('4 321' in mismatch_text, 'разница с разрядами')
    check('не является отчётностью' in mismatch_text, 'предупреждение, что файл не отчётность')
    check('200' in find_row(rows, 'Сверка выгрузок'), 'допуск сверки 1c показан')
    check(find_row(rows, 'Предупреждений').startswith('3'), 'количество предупреждений')

    print('\n3. Ключевые цифры отчёта на титульном листе')
    figures_ctx = make_context(
        balance_df=pd.DataFrame({
            'Актив/Пассив': ['А', 'А', 'П', 'П'],
            'Значение': [4_207_893.51, 3_818_268.92, -5_081_152.15, -2_945_010.28],
        }),
        pnl_df=pd.DataFrame({
            '1 уровень': ['Выручка', 'Себестоимость'],
            'Значение': [-1_359_246.45, 465_886.07],
        }),
    )
    rows = rc.build_cover_rows(figures_ctx, warnings=[])
    check(find_label(rows, 'КЛЮЧЕВЫЕ ЦИФРЫ ОТЧЁТА') == 'КЛЮЧЕВЫЕ ЦИФРЫ ОТЧЁТА', 'раздел ключевых цифр есть')
    check(find_row(rows, 'Актив (итог баланса)') == '8 026 162 тыс.ед.', 'актив на титульном листе')
    check(find_row(rows, 'Пассив (итог баланса)').startswith('-8 026 162'), 'пассив со знаком минус')
    check('СВЕДЕНО' in find_row(rows, 'Расхождение актив'), 'вердикт по сходимости баланса')
    check(find_row(rows, 'Выручка (ОПУ)') == '1 359 246 тыс.ед.', 'выручка по модулю')
    status_at = next(i for i, (label, _) in enumerate(rows) if label == 'СОСТОЯНИЕ ОТЧЁТА')
    figures_at = next(i for i, (label, _) in enumerate(rows) if label == 'КЛЮЧЕВЫЕ ЦИФРЫ ОТЧЁТА')
    modes_at = next(i for i, (label, _) in enumerate(rows) if label == 'РЕЖИМЫ СБОРКИ')
    check(status_at < figures_at < modes_at, 'цифры идут после статуса и до режимов')
    check(
        find_row(rows, 'Строк в расшифровке баланса') == '4'
        and find_row(rows, 'Строк в расшифровке ОПУ') == '2',
        'счётчики строк не задвоились в разделах',
    )
    no_figures = rc.build_cover_rows(
        make_context(balance_df=None, pnl_df=pd.DataFrame({'Значение': [1.0]})), warnings=[],
    )
    check('не собрана' in find_row(no_figures, 'Баланс'), 'без расшифровки баланса — не выдумываем цифры')
    check(find_row(no_figures, 'Выручка (ОПУ)') == '', 'без колонки уровней нет строки выручки')
    check(
        find_row(no_figures, 'Актив (итог баланса)') == '' and find_row(no_figures, 'Пассив (итог баланса)') == '',
        'при отсутствии баланса строк актива и пассива нет',
    )

    print('\n4. Формат допусков')
    check(format_tolerance_value('tolerance_rate_deviation', 0.3) == '30%', 'доля -> проценты')
    check(format_tolerance_value('nds_missing_values', 0.22) == '22%', 'ставка НДС -> проценты')
    check(format_tolerance_value('tolerance_balance', 5000) == '5 000 тыс.ед.', 'сумма -> разряды + единица')

    print('\n5. Запись в xlsx')
    out = tmp_dir / 'report.xlsx'
    with pd.ExcelWriter(out, engine='openpyxl') as writer:
        rc.write_cover_sheet(writer, rows)
        pd.DataFrame({'Колонка': [1, 2]}).to_excel(writer, sheet_name=rc.SHEET_BALANCE, index=False)

    wb = load_workbook(out)
    check(wb.sheetnames[0] == rc.COVER_SHEET, f'титульный лист первый: {wb.sheetnames}')
    ws = wb[rc.COVER_SHEET]
    check(ws.cell(1, 1).value == rc.COVER_LABEL_HEADER, 'шапка листа')
    check(ws.sheet_properties.tabColor.rgb.endswith(rc.TAB_COLOR_COVER), 'синий ярлычок титульного листа')
    check(ws.column_dimensions['A'].width == rc.COVER_LABEL_WIDTH, 'ширина колонки показателя')
    check(ws.column_dimensions['B'].width == rc.COVER_VALUE_WIDTH, 'ширина колонки значения')
    check(ws.cell(3, 2).alignment.wrap_text is True, 'перенос текста в значениях')

    section_row = next(r for r in range(2, ws.max_row + 1) if not ws.cell(r, 2).value)
    check(bool(ws.cell(section_row, 1).font.bold), f'заголовок раздела жирный (строка {section_row})')
    check(bool(ws.cell(section_row, 1).fill.start_color.rgb), 'заголовок раздела залит')

    print('\n6. Диагностика прогона')
    diagnostics = [label for label, _ in rows if '/' in label or label == 'sort_report.xlsx']
    check(not any('~$' in name for name in diagnostics), f'служебные ~$ исключены: {diagnostics}')
    check('sort_report.xlsx' in diagnostics, 'sort_report в списке диагностики')

    print('\n7. Сбой при сборе не ломает сохранение')
    check(rc.safe_build_cover_rows(object()) == [], 'битый контекст -> пустой лист, без исключения')

    print('\n8. Сводка предупреждений на титульном листе')
    warnings_list = ['Неизвестный контрагент: заменено на «3 лица» (×12)'] + [
        f'Предупреждение {i} (×{20 - i})' for i in range(2, 20)
    ]
    rows = rc.build_cover_rows(make_context(), warnings=warnings_list)
    check(find_row(rows, 'Предупреждений за прогон').startswith('19'),
          'всего предупреждений посчитано, не только показанных')
    check(find_row(rows, 'Неизвестный контрагент') == 'повторов: 12', 'текст с количеством повторов')
    check(find_row(rows, 'и ещё 7') == 'полный список — в логе прогона (app.log)',
          'остаток предупреждений указан одной строкой со ссылкой на лог')
    check(len([r for r in rows if r[1].startswith('повторов:')]) == rc.COVER_WARNINGS_LIMIT,
          f'показан лимит {rc.COVER_WARNINGS_LIMIT} строк')
    check(find_row(rows, 'Увязка ОПУ и баланса') != '', 'статус остался в разделе состояния')

    rows = rc.build_cover_rows(make_context(), warnings=[])
    check(find_row(rows, 'Предупреждений не было') == 'прогон прошёл без замечаний',
          'пустая сводка читается как «без замечаний»')

    rows = rc.build_cover_rows(make_context(), warnings=None)
    check('ПРЕДУПРЕЖДЕНИЯ ЗА ПРОГОН' not in [label for label, _ in rows],
          'без сводки раздел не создаётся (обратная совместимость)')

    print(f'\nSMOKE_OK ({PASSED} проверок)')

    try:
        for entry in sorted(tmp_dir.iterdir()):
            if entry.is_file():
                entry.unlink()
            else:
                for sub in entry.iterdir():
                    sub.unlink()
                entry.rmdir()
        tmp_dir.rmdir()
    except OSError:
        pass


if __name__ == '__main__':
    main()
